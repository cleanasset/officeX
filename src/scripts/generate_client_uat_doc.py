import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Reconfigure stdout for utf-8
sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    """Sets the background hex color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Sets inner padding for table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    """Adds a heading with consistent OfficeX corporate branding."""
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(4 if level == 1 else 3)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Calibri'
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(11, 31, 58) # Deep Navy
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(37, 99, 235) # Brand Blue
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    return h

def add_callout(doc, title, text, callout_type="tip"):
    """Adds a highlighted callout box with a colored border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "EFF6FF" if callout_type == "tip" else ("FEF2F2" if callout_type == "warning" else "F8FAFC")
    border_color = "3B82F6" if callout_type == "tip" else ("EF4444" if callout_type == "warning" else "64748B")
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
    
    # Left border styling
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_color)
    tcBorders.append(left)
    
    for side in ['top', 'bottom', 'right']:
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'none')
        tcBorders.append(edge)
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"{title}: ")
    r_title.font.name = 'Calibri'
    r_title.font.bold = True
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = RGBColor(11, 31, 58) if callout_type == "tip" else RGBColor(153, 27, 27)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    
    # Spacing after callout table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Building OfficeX End-to-End Client Testing & UAT Guide...")

doc = docx.Document()

# Configure standard 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Set base Normal style
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(10.5)
normal_style.font.color.rgb = RGBColor(51, 65, 85) # Slate 700

# ==============================================================================
# TITLE / COVER BLOCK
# ==============================================================================
p_meta = doc.add_paragraph()
p_meta.paragraph_format.space_before = Pt(0)
p_meta.paragraph_format.space_after = Pt(4)
r_badge = p_meta.add_run("OFFICEX COMMERCIAL REAL ESTATE ECOSYSTEM · SCALEZIX DELIVERY")
r_badge.font.size = Pt(9)
r_badge.font.bold = True
r_badge.font.color.rgb = RGBColor(37, 99, 235)

p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(4)
p_title.paragraph_format.space_after = Pt(6)
r_title = p_title.add_run("OfficeX Platform End-to-End Client Testing & UAT Guide")
r_title.font.size = Pt(24)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(11, 31, 58)

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_before = Pt(0)
p_sub.paragraph_format.space_after = Pt(16)
r_sub = p_sub.add_run("Comprehensive Operational Blueprint: Sign-In/Up, Rent Roll SaaS, Visitor Management, Statutory Compliance & Calendly Demo Flow")
r_sub.font.size = Pt(12)
r_sub.font.color.rgb = RGBColor(71, 85, 105)

# Document Control Table
doc_tbl = doc.add_table(rows=5, cols=2)
doc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
doc_tbl.autofit = False

doc_meta = [
    ("Document Control ID", "WO-OFFICEX-SCL-001 / UAT-REV-2.0"),
    ("Target Audience", "OfficeX Leadership, Clean Asset Services & Client Review Committee"),
    ("Implementation Partner", "Scalezix Engineering & Delivery Team"),
    ("Review Date & Version", "September 2026 · Version 1.2 (Production Baseline)"),
    ("System Core Directive", "Unified Multi-Tenant CRE SaaS Core ('Workspaces, Simplified')")
]

for idx, (label, val) in enumerate(doc_meta):
    c0 = doc_tbl.cell(idx, 0)
    c1 = doc_tbl.cell(idx, 1)
    c0.width = Inches(2.2)
    c1.width = Inches(4.3)
    set_cell_background(c0, "F1F5F9")
    set_cell_background(c1, "FFFFFF")
    set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
    set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(label)
    r0.font.bold = True
    r0.font.size = Pt(9.5)
    r0.font.color.rgb = RGBColor(15, 23, 42)
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(val)
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(51, 65, 85)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ==============================================================================
# SECTION 1: EXECUTIVE SUMMARY & ARCHITECTURAL FOUNDATION
# ==============================================================================
add_styled_heading(doc, "1. Executive Summary & Architectural Foundation", level=1)

p = doc.add_paragraph(
    "This document serves as the formal User Acceptance Testing (UAT) and functional handover guide "
    "for the OfficeX Commercial Real Estate (CRE) SaaS Platform. Developed strictly in accordance with "
    "the client's architectural directives ('Workspaces, Simplified'), this platform delivers an end-to-end "
    "operating system uniting Landlords, Property Managers, Corporate Tenants, Facility Management (FM) Operators, "
    "Leasing Brokers, and Service Vendors on a single, auditable core."
)
p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph(
    "Unlike legacy software that deploys isolated, fragmented tools, OfficeX operates as a unified multi-tenant architecture "
    "sharing a central organizational master: Organization → Properties → Leases → Visitors → Compliance → Audits → Financials. "
    "All sandbox pills, hardcoded mock hints, and static demo codes have been completely decommissioned. The entire system is "
    "now powered by real cryptographic verification, free standard SMTP email dispatch, live financial algorithms, and dynamic workspace routing."
)
p.paragraph_format.space_after = Pt(10)

add_callout(
    doc,
    "Client Guiding Principle",
    "OfficeX sign-in is an enterprise-grade front door reflecting 'Workspaces, Simplified' rather than an internal module picker. "
    "Users authenticate once via their work identity and are dynamically routed to their active workspace based on their verified organization roles.",
    callout_type="tip"
)

# ==============================================================================
# SECTION 2: THE 9 UNIFIED WORKSPACES & PORTALS DIRECTORY
# ==============================================================================
add_styled_heading(doc, "2. The 9 Unified Workspaces & Portals Directory", level=1)

p = doc.add_paragraph(
    "The OfficeX platform consists of 9 integrated workspaces accessible across public marketing and auth-gated enterprise environments. "
    "Clients and testers can navigate directly to these environments using the canonical routes below:"
)
p.paragraph_format.space_after = Pt(8)

# 9 Portals Table
portals_tbl = doc.add_table(rows=10, cols=4)
portals_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
portals_tbl.autofit = False

headers = ["Workspace / Portal", "Primary Role & Persona", "Route / Direct URL", "Core Operational Focus"]
widths = [Inches(1.5), Inches(1.5), Inches(1.5), Inches(2.0)]

for c_idx, h_text in enumerate(headers):
    cell = portals_tbl.cell(0, c_idx)
    cell.width = widths[c_idx]
    set_cell_background(cell, "0B1F3A")
    set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(255, 255, 255)

portals_data = [
    ("1. Public Discovery & CRE Marketplace", "Prospective Tenants & Public Investors", "/\n/marketplace\n/public/search", "Verified Grade-A building discovery, map search, 3D tour requests, site visit booking via Calendly."),
    ("2. Commercial Landlord Desk", "Property Owners & Asset Directors (Ravi Singhal)", "/properties\n/operate/rent-roll", "Master portfolio oversight, rent roll cash flow, vacancy loss, lease escalations, and NOI yield tracking."),
    ("3. FM Command Centre", "Facility Operations Managers (Sunil Verma)", "/ops\n/operate/ppm\n/operate/helpdesk", "Building operations, work orders, reactive tickets, Planned Preventive Maintenance (PPM), and SLA escalations."),
    ("4. FM Marketplace", "Corporate Tenants & Building Owners", "/fm-marketplace", "On-demand facility service procurement: Deep cleaning, HVAC maintenance, pest control, and interior fitouts."),
    ("5. Corporate Tenant Portal", "Workplace Admin & Facility Leads (Pooja Iyer)", "/tenant", "Employee desk booking, office access passes, service requests, utility consumption, and rent invoice review."),
    ("6. Leasing Broker CRM", "Agency Brokers & Channel Partners (Arjun Kapoor)", "/leasing\n/leasing/pipeline", "Listing distribution, tenant lead CRM, site visit schedules, commission tracking, and digital LOI generation."),
    ("7. Service Vendor Hub", "Contractors & Facility Vendors (Karan Mehra)", "/vendor", "Work order dispatch acceptance, technician scheduling, milestone billing, and AMC contract compliance."),
    ("8. Super Admin Console", "Platform Governance & Security (Deepak Sharma)", "/admin", "Organization master, KYC verification, role-based access control (RBAC), security audit logs, and privileged MFA."),
    ("9. Executive MIS & Analytics", "Institutional Investors & Fund Managers", "/reporting\n/operate/compliance", "Portfolio-wide ESG benchmarks, statutory compliance audits, tenant aging, and gross-to-net yield analytics.")
]

for row_idx, data in enumerate(portals_data, start=1):
    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
    for col_idx, text in enumerate(data):
        cell = portals_tbl.cell(row_idx, col_idx)
        cell.width = widths[col_idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif col_idx == 2:
            r.font.name = 'Consolas'
            r.font.color.rgb = RGBColor(37, 99, 235)
        else:
            r.font.color.rgb = RGBColor(71, 85, 105)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# ==============================================================================
# SECTION 3: REAL SIGN-IN & SIGN-UP FLOW WALKTHROUGH
# ==============================================================================
add_styled_heading(doc, "3. Module 1: Real Sign-In & Sign-Up Flow Walkthrough", level=1)

p = doc.add_paragraph(
    "The authentication layer has been re-architected into a clean, identifier-first experience backed by "
    "real SMTP email dispatch and cryptographic verification. There are zero simulated sandbox pills."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "A. Role-Based Self-Registration Flow (/signup)", level=2)
p = doc.add_paragraph(
    "1. Navigate to: http://localhost:3000/signup\n"
    "2. Enter Full Name (e.g., 'Vikram Malhotra') and Corporate Work Email (e.g., 'vikram@horizonrealty.com').\n"
    "3. Enter a 10-Digit Mobile Number (e.g., '9820112233'). Country prefix (+91) is automatically applied.\n"
    "4. Choose Intended Workspace Role:\n"
    "   • Property Owner / Asset Manager\n"
    "   • Commercial Leasing Broker\n"
    "   • Facility / Service Vendor\n"
    "5. Set a Secure Password (minimum 8 characters with uppercase, numbers, and symbols).\n"
    "6. Click 'Create Account & Verify'.\n"
    "7. The system dynamically generates a cryptographically random 6-digit PIN and immediately dispatches it via free SMTP to the entered email.\n"
    "8. Enter the 6-digit code received in the inbox. Upon successful verification, the account is provisioned in the database and redirected to the workspace."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "B. Smart Identifier-First Sign-In Flow (/login)", level=2)
p = doc.add_paragraph(
    "1. Navigate to: http://localhost:3000/login\n"
    "2. Smart Identifier Detection:\n"
    "   • The user enters their Work Email (e.g., 'ravi@acme.com') OR Mobile Number (e.g., '9876543210').\n"
    "   • The system instantly detects the format without requiring the user to specify whether it is an email or phone.\n"
    "3. Authentication Pathways:\n"
    "   • Password Authentication: Enter password with 1-click show/hide toggle.\n"
    "   • One-Time Passcode (OTP): Click 'Email me a one-time code instead' to dispatch a 6-digit PIN.\n"
    "   • Enterprise SSO: If entering a corporate domain (e.g. '@dlf.in' or '@tcs.com'), the form routes to Okta/Azure AD.\n"
    "   • Privileged MFA: Super Administrator ('admin@officex.in') automatically triggers mandatory Level-2 MFA.\n"
    "4. Multi-Workspace Context Chooser ('Where would you like to work today?'):\n"
    "   • If a user holds multiple roles across organizations (e.g., Landlord of Acme Realty AND Executive at NovaTech),\n"
    "     OfficeX presents a branded context selector modal displaying badges, properties, and last-used workspace.\n"
    "   • Selecting a card routes the user directly to the appropriate portal without requiring re-authentication."
)
p.paragraph_format.space_after = Pt(10)

add_callout(
    doc,
    "Anti-Brute Force Security",
    "The OTP engine enforces a strict 30-second cooldown between resend requests and permanently locks a code after 5 consecutive incorrect attempts. "
    "Codes automatically expire in 10 minutes and are consumed upon first use to eliminate replay vulnerabilities.",
    callout_type="tip"
)

# ==============================================================================
# SECTION 4: RENT ROLL SAAS MODULE WALKTHROUGH
# ==============================================================================
add_styled_heading(doc, "4. Module 2: Rent Roll SaaS Operating System Walkthrough", level=1)

p = doc.add_paragraph(
    "The OfficeX Rent Roll SaaS module transforms traditional static spreadsheets into a real-time, auditable "
    "operating system. It serves as the single source of truth for commercial leases, rental escalations, CAM billing, "
    "and occupancy forecasting across commercial office towers."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "A. How a Building Owner or Property Manager Operates (/operate/rent-roll)", level=2)
p = doc.add_paragraph(
    "1. Navigate to: http://localhost:3000/operate/rent-roll (or via Landlord Desk at /properties)\n"
    "2. Select Active Commercial Asset:\n"
    "   • Choose from the property dropdown (e.g., 'Apex Horizon Tower - Bandra Kurla Complex, Mumbai' or 'GIFT One Tower, GIFT City').\n"
    "3. Executive KPI Strip (Live Calculations):\n"
    "   • Total Leasable Area: 285,000 Sq. Ft.\n"
    "   • Current Occupancy Rate: 92.4% (Occupied: 263,340 Sq. Ft. · Vacant: 21,660 Sq. Ft.)\n"
    "   • Monthly Gross Rent Collection: ₹4.82 Cr / month\n"
    "   • Weighted Average Lease Expiry (WALE): 3.8 Years\n"
    "   • Monthly Vacancy Loss: ₹38.98 Lakhs / month (calculated from unleased space at market rental)\n"
    "4. Unit-by-Unit Master Rent Roll Table:\n"
    "   • Unit / Suite ID (e.g., 'Suite 401 - Wing A')\n"
    "   • Tenant Entity Name (e.g., 'Morgan Stanley India Support', 'KPMG Shared Services')\n"
    "   • Leased Area (Sq. Ft.) and Fitout Status (Fully Furnished, Bare Shell, Warm Shell)\n"
    "   • Base Rent (₹ / Sq. Ft. / Month) and Common Area Maintenance (CAM ₹ / Sq. Ft.)\n"
    "   • Current Monthly Rent and Escalation Schedule (e.g., '15% every 36 months' or '5% annual')\n"
    "   • Security Deposit Held in Escrow (Standard 6 months gross rent)\n"
    "   • Lease Expiry Date and Active Expiry Alert Badge (Green > 180 days, Amber 90 days, Red < 30 days)\n"
    "5. Actionable Operational Buttons:\n"
    "   • 'Add New Lease': Opens digital lease creation form with auto-escalation compounding.\n"
    "   • 'Upload Master Sheet': Import existing client Excel rent rolls directly into PostgreSQL.\n"
    "   • 'Export Rent Roll': Download audit-ready Excel/CSV sheet with formula preservation.\n"
    "   • '3-Year Revenue Forecast': Visualizes compounded revenue growth factoring in contracted escalations."
)
p.paragraph_format.space_after = Pt(10)

# ==============================================================================
# SECTION 5: VISITOR MANAGEMENT SAAS WALKTHROUGH
# ==============================================================================
add_styled_heading(doc, "5. Module 3: Enterprise Visitor Management SaaS Walkthrough", level=1)

p = doc.add_paragraph(
    "Designed for Grade-A commercial office buildings, the Visitor Management module eliminates paper logbooks, "
    "accelerates reception throughput, and maintains a strict real-time audit trail of all guests, contractors, and VIPs."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "A. End-to-End Visitor Journey (/operate/visitors)", level=2)
p = doc.add_paragraph(
    "1. Navigate to: http://localhost:3000/operate/visitors (or via FM Command Centre at /ops)\n"
    "2. Pathway A: Pre-Registration (Host Invites Guest):\n"
    "   • Tenant host inputs visitor name, company, mobile number, visit date, and meeting room.\n"
    "   • The platform automatically generates a Digital QR Pass and dispatches it to the visitor via WhatsApp/Email.\n"
    "   • Upon arrival at the turnstile or reception, the visitor scans their QR pass for 3-second check-in.\n"
    "3. Pathway B: Reception Walk-In Kiosk:\n"
    "   • Unscheduled visitor approaches the front desk or self-service tablet.\n"
    "   • Visitor enters mobile number, captures webcam photo, and accepts building safety/NDA policies.\n"
    "   • System verifies visitor via mobile OTP and generates a visitor pass with thermal badge printing.\n"
    "4. Instant Host Notification:\n"
    "   • The host tenant receives an automated alert: 'Your visitor [Guest Name] has arrived at Tower Reception.'\n"
    "5. Real-Time Security Command & Evacuation:\n"
    "   • Security desk dashboard displays live occupancy: Total Checked-In, Overstay Alerts (>8 hours), and Check-Outs.\n"
    "   • In the event of a fire drill or emergency, 1-click 'Evacuation Roll-Call' generates an instant roster of all guests currently on-site."
)
p.paragraph_format.space_after = Pt(10)

# ==============================================================================
# SECTION 6: STATUTORY COMPLIANCE & SAFETY MANAGEMENT WALKTHROUGH
# ==============================================================================
add_styled_heading(doc, "6. Module 4: Statutory Compliance & Safety Management Walkthrough", level=1)

p = doc.add_paragraph(
    "Commercial office operations in India require strict adherence to local municipal, fire, environmental, and "
    "labor regulations. The Compliance Management SaaS module monitors certificate expirations, audits, and vendor AMCs."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "A. Regulatory Audit & Expiry Engine (/operate/compliance)", level=2)
p = doc.add_paragraph(
    "1. Navigate to: http://localhost:3000/operate/compliance\n"
    "2. The 4 Mandatory Commercial Real Estate Regulatory Pillars:\n"
    "   • Pillar 1: Fire Safety NOC (Chief Fire Officer annual inspection & hydrostatic test certificates).\n"
    "   • Pillar 2: Lift & Escalator PWD Safety License (Annual government inspectorate approval & rope certificates).\n"
    "   • Pillar 3: Diesel Generator (DG Set) Consent to Operate (State Pollution Control Board emission norms).\n"
    "   • Pillar 4: Structural Stability & Building Occupancy Certificate (Municipal Corporation sanctioned plan).\n"
    "3. Smart Expiry Tracker & Escalation Matrix:\n"
    "   • Green (>90 Days Remaining): Compliant / Normal state.\n"
    "   • Amber (30 to 90 Days Remaining): Renewal triggered; notification sent to Facility Manager.\n"
    "   • Red (<30 Days Remaining): Critical warning; escalated to Vice President of Property Operations.\n"
    "4. Document Vault & Evidence Repository:\n"
    "   • Secure storage for scanned government licenses, test reports, challans, and approved drawings.\n"
    "   • Third-party auditor access allows read-only verification during investor due diligence."
)
p.paragraph_format.space_after = Pt(10)

# ==============================================================================
# SECTION 7: CALENDLY / CALENDAR SCHEDULING INTEGRATION
# ==============================================================================
add_styled_heading(doc, "7. Module 5: Calendly & Calendar Scheduling Integration", level=1)

p = doc.add_paragraph(
    "To drive maximum lead conversion without creating friction, OfficeX integrates direct meeting scheduling "
    "seamlessly into the application flow."
)
p.paragraph_format.space_after = Pt(6)

add_styled_heading(doc, "A. Direct Booking Flow (30-Minute Consultation)", level=2)
p = doc.add_paragraph(
    "1. Production Booking Link: https://calendly.com/admin-cleanasset/30min\n"
    "2. In-App Trigger Points:\n"
    "   • Marketing Header: 'Schedule Tour' / 'Book a Demo' navigation link.\n"
    "   • Property Listing Pages: 'Schedule Inspection' / 'Book Site Visit' buttons on individual towers.\n"
    "   • SaaS Solution Banners: 'Talk to Real Estate Advisor' across Landlord and FM portals.\n"
    "3. High-Performance Modal Experience:\n"
    "   • Clicking the trigger immediately opens a responsive, branded popup modal within OfficeX.\n"
    "   • The user selects their preferred date, time zone, and meeting time directly on screen without leaving the platform.\n"
    "   • Automated Calendar Sync: Dispatches immediate Google Calendar / Microsoft Outlook invites with Google Meet video links "
    "     to both the client prospect and the OfficeX / Clean Asset advisory team."
)
p.paragraph_format.space_after = Pt(10)

# ==============================================================================
# SECTION 8: STEP-BY-STEP CLIENT UAT TEST SCRIPTS
# ==============================================================================
add_styled_heading(doc, "8. Step-by-Step Client UAT Test Scripts", level=1)

p = doc.add_paragraph(
    "The following test cases provide an actionable, click-by-click testing procedure for client stakeholders and QA teams. "
    "Execute these tests in sequence to validate the entire platform flow:"
)
p.paragraph_format.space_after = Pt(8)

# UAT Table
uat_tbl = doc.add_table(rows=11, cols=4)
uat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
uat_tbl.autofit = False

uat_headers = ["Test ID & Persona", "Module / URL", "Step-by-Step Actions & Input Data", "Expected Verification Result"]
uat_widths = [Inches(1.2), Inches(1.3), Inches(2.2), Inches(1.8)]

for c_idx, h_text in enumerate(uat_headers):
    cell = uat_tbl.cell(0, c_idx)
    cell.width = uat_widths[c_idx]
    set_cell_background(cell, "0B1F3A")
    set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(255, 255, 255)

uat_data = [
    ("TC-01\nNew User", "/signup", "1. Fill Name, Work Email, Mobile (+91).\n2. Select 'Property Owner' role.\n3. Enter password & click 'Create Account'.", "Real 6-digit OTP sent via SMTP to inbox. Entering code creates account with zero sandbox pills."),
    ("TC-02\nLandlord", "/login", "1. Enter 'ravi@acme.com'.\n2. System detects email & shows password box.\n3. Enter 'OfficeX@2026' & Submit.", "Authenticates cleanly and displays post-auth workspace chooser ('Where would you like to work today?')."),
    ("TC-03\nMobile User", "/login", "1. Enter '9876543210'.\n2. System detects phone & routes to OTP.\n3. Request verification code.", "Dispatches real 6-digit code with 30s timer. Entering code routes to FM Command Centre (/ops)."),
    ("TC-04\nAsset Owner", "/operate/rent-roll", "1. Select 'Apex Horizon Tower'.\n2. Review Gross Rent, WALE, and Vacancy.\n3. Click 'Add New Lease'.", "KPIs compute dynamically. New lease modal allows configuring sq. ft., base rent, CAM, and escalation %."),
    ("TC-05\nAuditor", "/operate/rent-roll", "1. Click 'Export Rent Roll' button at top right.", "System immediately downloads audit-ready spreadsheet with tenant breakdown and escalations preserved."),
    ("TC-06\nHost Tenant", "/operate/visitors", "1. Click 'Pre-Register Guest'.\n2. Enter Name, Mobile, Date, and Meeting Room.", "System generates unique Digital QR Pass with 1-click WhatsApp/Email dispatch link."),
    ("TC-07\nSecurity Desk", "/operate/visitors", "1. Search visitor by mobile or scan pass.\n2. Click 'Check In'.", "Visitor marked 'Checked-In', live building occupancy increments, and host receives instant notification."),
    ("TC-08\nCompliance Head", "/operate/compliance", "1. Review 4 statutory categories.\n2. Click 'Fire Safety NOC' card.", "Shows validity dates, attached inspection certificate, renewal checklist, and 90-day warning badge."),
    ("TC-09\nProspect", "/ (Homepage)\n/marketplace", "1. Click 'Book a Demo' or 'Schedule Tour'.", "Calendly modal opens smoothly with https://calendly.com/admin-cleanasset/30min. Zero page reload."),
    ("TC-10\nSuper Admin", "/login", "1. Enter 'admin@officex.in'.\n2. Submit credentials.", "Enforces mandatory Level-2 MFA before granting access to the Super Admin Governance Console (/admin).")
]

for row_idx, data in enumerate(uat_data, start=1):
    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
    for col_idx, text in enumerate(data):
        cell = uat_tbl.cell(row_idx, col_idx)
        cell.width = uat_widths[col_idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif col_idx == 1:
            r.font.name = 'Consolas'
            r.font.color.rgb = RGBColor(37, 99, 235)
        else:
            r.font.color.rgb = RGBColor(71, 85, 105)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# ==============================================================================
# SECTION 9: TECHNICAL HANDOVER & ENVIRONMENT SETTINGS
# ==============================================================================
add_styled_heading(doc, "9. Technical Handover & Production Configuration", level=1)

p = doc.add_paragraph(
    "For the client's IT and technical operations team, the following environment variables in `.env.local` "
    "govern core platform integrations. All services operate under 100% free production-ready tiers:"
)
p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph(
    "• SMTP_HOST: Standard SMTP server address (e.g. 'smtp.gmail.com' or 'smtp-relay.brevo.com').\n"
    "• SMTP_PORT: Secure port ('465' for SSL or '587' for TLS).\n"
    "• SMTP_USER: Outgoing sender mailbox address.\n"
    "• SMTP_PASS: 16-character application-specific password (Google App Password / Brevo API key).\n"
    "• SMTP_FROM: Branded sender header (e.g. 'OfficeX Security <noreply@officex.in>').\n"
    "• NEXT_PUBLIC_SUPABASE_URL & ANON_KEY: High-performance PostgreSQL database backend."
)
p.paragraph_format.space_after = Pt(10)

p_sign = doc.add_paragraph()
p_sign.paragraph_format.space_before = Pt(16)
p_sign.paragraph_format.space_after = Pt(2)
r_sign = p_sign.add_run("Delivered & Certified by Scalezix Engineering & Delivery Team")
r_sign.font.bold = True
r_sign.font.size = Pt(11)
r_sign.font.color.rgb = RGBColor(11, 31, 58)

p_foot = doc.add_paragraph()
p_foot.paragraph_format.space_after = Pt(0)
r_foot = p_foot.add_run("OfficeX Platform Ecosystem · Workspaces, Simplified · Confidential Document for Client UAT")
r_foot.font.size = Pt(9)
r_foot.font.italic = True
r_foot.font.color.rgb = RGBColor(100, 116, 139)

# Save Word document
docx_path = os.path.abspath("client response/OfficeX_End_To_End_Client_Testing_Guide.docx")
doc.save(docx_path)
print(f"✅ Successfully generated Word Document: {docx_path}")

# ==============================================================================
# GENERATE COMPANION MARKDOWN DOCUMENT
# ==============================================================================
md_content = """# OfficeX Platform End-to-End Client Testing & UAT Guide
**Comprehensive Operational Blueprint: Sign-In/Up, Rent Roll SaaS, Visitor Management, Statutory Compliance & Calendly Demo Flow**

* **Document ID:** `WO-OFFICEX-SCL-001 / UAT-REV-2.0`
* **Target Audience:** OfficeX Leadership, Clean Asset Services & Client Review Committee
* **Implementation Partner:** Scalezix Engineering & Delivery Team
* **Review Date & Version:** September 2026 · Version 1.2 (Production Baseline)
* **System Core Directive:** Unified Multi-Tenant CRE SaaS Core (*"Workspaces, Simplified"*)

---

## 1. Executive Summary & Architectural Foundation

This document serves as the formal User Acceptance Testing (UAT) and functional handover guide for the OfficeX Commercial Real Estate (CRE) SaaS Platform. Developed strictly in accordance with the client's architectural directives (*"Workspaces, Simplified"*), this platform delivers an end-to-end operating system uniting Landlords, Property Managers, Corporate Tenants, Facility Management (FM) Operators, Leasing Brokers, and Service Vendors on a single, auditable core.

Unlike legacy software that deploys isolated, fragmented tools, OfficeX operates as a unified multi-tenant architecture sharing a central organizational master: **Organization → Properties → Leases → Visitors → Compliance → Audits → Financials**. All sandbox pills, hardcoded mock hints, and static demo codes have been completely decommissioned. The entire system is now powered by real cryptographic verification, free standard SMTP email dispatch, live financial algorithms, and dynamic workspace routing.

> **Client Guiding Principle:**
> OfficeX sign-in is an enterprise-grade front door reflecting *"Workspaces, Simplified"* rather than an internal module picker. Users authenticate once via their work identity and are dynamically routed to their active workspace based on their verified organization roles.

---

## 2. The 9 Unified Workspaces & Portals Directory

The OfficeX platform consists of 9 integrated workspaces accessible across public marketing and auth-gated enterprise environments:

| Workspace / Portal | Primary Role & Persona | Route / Direct URL | Core Operational Focus |
| :--- | :--- | :--- | :--- |
| **1. Public Discovery & CRE Marketplace** | Prospective Tenants & Public Investors | `/`<br>`/marketplace`<br>`/public/search` | Verified Grade-A building discovery, map search, 3D tour requests, site visit booking via Calendly. |
| **2. Commercial Landlord Desk** | Property Owners & Asset Directors (*Ravi Singhal*) | `/properties`<br>`/operate/rent-roll` | Master portfolio oversight, rent roll cash flow, vacancy loss, lease escalations, and NOI yield tracking. |
| **3. FM Command Centre** | Facility Operations Managers (*Sunil Verma*) | `/ops`<br>`/operate/ppm`<br>`/operate/helpdesk` | Building operations, work orders, reactive tickets, Planned Preventive Maintenance (PPM), and SLA escalations. |
| **4. FM Marketplace** | Corporate Tenants & Building Owners | `/fm-marketplace` | On-demand facility service procurement: Deep cleaning, HVAC maintenance, pest control, and interior fitouts. |
| **5. Corporate Tenant Portal** | Workplace Admin & Facility Leads (*Pooja Iyer*) | `/tenant` | Employee desk booking, office access passes, service requests, utility consumption, and rent invoice review. |
| **6. Leasing Broker CRM** | Agency Brokers & Channel Partners (*Arjun Kapoor*) | `/leasing`<br>`/leasing/pipeline` | Listing distribution, tenant lead CRM, site visit schedules, commission tracking, and digital LOI generation. |
| **7. Service Vendor Hub** | Contractors & Facility Vendors (*Karan Mehra*) | `/vendor` | Work order dispatch acceptance, technician scheduling, milestone billing, and AMC contract compliance. |
| **8. Super Admin Console** | Platform Governance & Security (*Deepak Sharma*) | `/admin` | Organization master, KYC verification, role-based access control (RBAC), security audit logs, and privileged MFA. |
| **9. Executive MIS & Analytics** | Institutional Investors & Fund Managers | `/reporting`<br>`/operate/compliance` | Portfolio-wide ESG benchmarks, statutory compliance audits, tenant aging, and gross-to-net yield analytics. |

---

## 3. Module 1: Real Sign-In & Sign-Up Flow Walkthrough

The authentication layer has been re-architected into a clean, identifier-first experience backed by real SMTP email dispatch and cryptographic verification. There are zero simulated sandbox pills.

### A. Role-Based Self-Registration Flow (`/signup`)
1. **Navigate to:** `http://localhost:3000/signup`
2. **Enter Full Name** (e.g., `"Vikram Malhotra"`) and **Corporate Work Email** (e.g., `"vikram@horizonrealty.com"`).
3. **Enter a 10-Digit Mobile Number** (e.g., `"9820112233"`). Country prefix (`+91`) is automatically applied.
4. **Choose Intended Workspace Role:**
   * Property Owner / Asset Manager
   * Commercial Leasing Broker
   * Facility / Service Vendor
5. **Set a Secure Password** (minimum 8 characters with uppercase, numbers, and symbols).
6. Click **Create Account & Verify**.
7. The system dynamically generates a cryptographically random 6-digit PIN and immediately dispatches it via free SMTP to the entered email.
8. Enter the 6-digit code received in the inbox. Upon successful verification, the account is provisioned in the database and redirected to the workspace.

### B. Smart Identifier-First Sign-In Flow (`/login`)
1. **Navigate to:** `http://localhost:3000/login`
2. **Smart Identifier Detection:**
   * The user enters their Work Email (e.g., `"ravi@acme.com"`) OR Mobile Number (e.g., `"9876543210"`).
   * The system instantly detects the format without requiring the user to specify whether it is an email or phone.
3. **Authentication Pathways:**
   * **Password Authentication:** Enter password with 1-click show/hide toggle.
   * **One-Time Passcode (OTP):** Click *"Email me a one-time code instead"* to dispatch a 6-digit PIN.
   * **Enterprise SSO:** If entering a corporate domain (e.g. `@dlf.in` or `@tcs.com`), the form routes to Okta/Azure AD.
   * **Privileged MFA:** Super Administrator (`admin@officex.in`) automatically triggers mandatory Level-2 MFA.
4. **Multi-Workspace Context Chooser (*"Where would you like to work today?"*):**
   * If a user holds multiple roles across organizations (e.g., Landlord of Acme Realty AND Executive at NovaTech), OfficeX presents a branded context selector modal displaying badges, properties, and last-used workspace.
   * Selecting a card routes the user directly to the appropriate portal without requiring re-authentication.

> **Anti-Brute Force Security:**
> The OTP engine enforces a strict 30-second cooldown between resend requests and permanently locks a code after 5 consecutive incorrect attempts. Codes automatically expire in 10 minutes and are consumed upon first use to eliminate replay vulnerabilities.

---

## 4. Module 2: Rent Roll SaaS Operating System Walkthrough

The OfficeX Rent Roll SaaS module transforms traditional static spreadsheets into a real-time, auditable operating system. It serves as the single source of truth for commercial leases, rental escalations, CAM billing, and occupancy forecasting across commercial office towers.

### A. How a Building Owner or Property Manager Operates (`/operate/rent-roll`)
1. **Navigate to:** `http://localhost:3000/operate/rent-roll` (or via Landlord Desk at `/properties`)
2. **Select Active Commercial Asset:**
   * Choose from the property dropdown (e.g., *"Apex Horizon Tower - Bandra Kurla Complex, Mumbai"* or *"GIFT One Tower, GIFT City"*).
3. **Executive KPI Strip (Live Calculations):**
   * **Total Leasable Area:** 285,000 Sq. Ft.
   * **Current Occupancy Rate:** 92.4% (Occupied: 263,340 Sq. Ft. · Vacant: 21,660 Sq. Ft.)
   * **Monthly Gross Rent Collection:** ₹4.82 Cr / month
   * **Weighted Average Lease Expiry (WALE):** 3.8 Years
   * **Monthly Vacancy Loss:** ₹38.98 Lakhs / month (calculated from unleased space at market rental)
4. **Unit-by-Unit Master Rent Roll Table:**
   * **Unit / Suite ID** (e.g., `"Suite 401 - Wing A"`)
   * **Tenant Entity Name** (e.g., `"Morgan Stanley India Support"`, `"KPMG Shared Services"`)
   * **Leased Area (Sq. Ft.)** and **Fitout Status** (Fully Furnished, Bare Shell, Warm Shell)
   * **Base Rent** (₹ / Sq. Ft. / Month) and **Common Area Maintenance (CAM)** (₹ / Sq. Ft.)
   * **Current Monthly Rent** and **Escalation Schedule** (e.g., `"15% every 36 months"` or `"5% annual"`)
   * **Security Deposit** Held in Escrow (Standard 6 months gross rent)
   * **Lease Expiry Date** and **Active Expiry Alert Badge** (Green > 180 days, Amber 90 days, Red < 30 days)
5. **Actionable Operational Buttons:**
   * **"Add New Lease":** Opens digital lease creation form with auto-escalation compounding.
   * **"Upload Master Sheet":** Import existing client Excel rent rolls directly into PostgreSQL.
   * **"Export Rent Roll":** Download audit-ready Excel/CSV sheet with formula preservation.
   * **"3-Year Revenue Forecast":** Visualizes compounded revenue growth factoring in contracted escalations.

---

## 5. Module 3: Enterprise Visitor Management SaaS Walkthrough

Designed for Grade-A commercial office buildings, the Visitor Management module eliminates paper logbooks, accelerates reception throughput, and maintains a strict real-time audit trail of all guests, contractors, and VIPs.

### A. End-to-End Visitor Journey (`/operate/visitors`)
1. **Navigate to:** `http://localhost:3000/operate/visitors` (or via FM Command Centre at `/ops`)
2. **Pathway A: Pre-Registration (Host Invites Guest):**
   * Tenant host inputs visitor name, company, mobile number, visit date, and meeting room.
   * The platform automatically generates a Digital QR Pass and dispatches it to the visitor via WhatsApp/Email.
   * Upon arrival at the turnstile or reception, the visitor scans their QR pass for 3-second check-in.
3. **Pathway B: Reception Walk-In Kiosk:**
   * Unscheduled visitor approaches the front desk or self-service tablet.
   * Visitor enters mobile number, captures webcam photo, and accepts building safety/NDA policies.
   * System verifies visitor via mobile OTP and generates a visitor pass with thermal badge printing.
4. **Instant Host Notification:**
   * The host tenant receives an automated alert: *"Your visitor [Guest Name] has arrived at Tower Reception."*
5. **Real-Time Security Command & Evacuation:**
   * Security desk dashboard displays live occupancy: Total Checked-In, Overstay Alerts (>8 hours), and Check-Outs.
   * In the event of a fire drill or emergency, 1-click **"Evacuation Roll-Call"** generates an instant roster of all guests currently on-site.

---

## 6. Module 4: Statutory Compliance & Safety Management Walkthrough

Commercial office operations in India require strict adherence to local municipal, fire, environmental, and labor regulations. The Compliance Management SaaS module monitors certificate expirations, audits, and vendor AMCs.

### A. Regulatory Audit & Expiry Engine (`/operate/compliance`)
1. **Navigate to:** `http://localhost:3000/operate/compliance`
2. **The 4 Mandatory Commercial Real Estate Regulatory Pillars:**
   * **Pillar 1: Fire Safety NOC** (Chief Fire Officer annual inspection & hydrostatic test certificates).
   * **Pillar 2: Lift & Escalator PWD Safety License** (Annual government inspectorate approval & rope certificates).
   * **Pillar 3: Diesel Generator (DG Set) Consent to Operate** (State Pollution Control Board emission norms).
   * **Pillar 4: Structural Stability & Building Occupancy Certificate** (Municipal Corporation sanctioned plan).
3. **Smart Expiry Tracker & Escalation Matrix:**
   * **Green (>90 Days Remaining):** Compliant / Normal state.
   * **Amber (30 to 90 Days Remaining):** Renewal triggered; notification sent to Facility Manager.
   * **Red (<30 Days Remaining):** Critical warning; escalated to Vice President of Property Operations.
4. **Document Vault & Evidence Repository:**
   * Secure storage for scanned government licenses, test reports, challans, and approved drawings.
   * Third-party auditor access allows read-only verification during investor due diligence.

---

## 7. Module 5: Calendly & Calendar Scheduling Integration

To drive maximum lead conversion without creating friction, OfficeX integrates direct meeting scheduling seamlessly into the application flow.

### A. Direct Booking Flow (30-Minute Consultation)
1. **Production Booking Link:** `https://calendly.com/admin-cleanasset/30min`
2. **In-App Trigger Points:**
   * **Marketing Header:** *"Schedule Tour"* / *"Book a Demo"* navigation link.
   * **Property Listing Pages:** *"Schedule Inspection"* / *"Book Site Visit"* buttons on individual towers.
   * **SaaS Solution Banners:** *"Talk to Real Estate Advisor"* across Landlord and FM portals.
3. **High-Performance Modal Experience:**
   * Clicking the trigger immediately opens a responsive, branded popup modal within OfficeX.
   * The user selects their preferred date, time zone, and meeting time directly on screen without leaving the platform.
   * **Automated Calendar Sync:** Dispatches immediate Google Calendar / Microsoft Outlook invites with Google Meet video links to both the client prospect and the OfficeX / Clean Asset advisory team.

---

## 8. Step-by-Step Client UAT Test Scripts

Execute these test cases in sequence to validate the entire platform flow:

| Test ID & Persona | Module / URL | Step-by-Step Actions & Input Data | Expected Verification Result |
| :--- | :--- | :--- | :--- |
| **TC-01**<br>New User | `/signup` | 1. Fill Name, Work Email, Mobile (+91).<br>2. Select 'Property Owner' role.<br>3. Enter password & click 'Create Account'. | Real 6-digit OTP sent via SMTP to inbox. Entering code creates account with zero sandbox pills. |
| **TC-02**<br>Landlord | `/login` | 1. Enter `ravi@acme.com`.<br>2. System detects email & shows password box.<br>3. Enter `OfficeX@2026` & Submit. | Authenticates cleanly and displays post-auth workspace chooser (*'Where would you like to work today?'*). |
| **TC-03**<br>Mobile User | `/login` | 1. Enter `9876543210`.<br>2. System detects phone & routes to OTP.<br>3. Request verification code. | Dispatches real 6-digit code with 30s timer. Entering code routes to FM Command Centre (`/ops`). |
| **TC-04**<br>Asset Owner | `/operate/rent-roll` | 1. Select 'Apex Horizon Tower'.<br>2. Review Gross Rent, WALE, and Vacancy.<br>3. Click 'Add New Lease'. | KPIs compute dynamically. New lease modal allows configuring sq. ft., base rent, CAM, and escalation %. |
| **TC-05**<br>Auditor | `/operate/rent-roll` | 1. Click 'Export Rent Roll' button at top right. | System immediately downloads audit-ready spreadsheet with tenant breakdown and escalations preserved. |
| **TC-06**<br>Host Tenant | `/operate/visitors` | 1. Click 'Pre-Register Guest'.<br>2. Enter Name, Mobile, Date, and Meeting Room. | System generates unique Digital QR Pass with 1-click WhatsApp/Email dispatch link. |
| **TC-07**<br>Security Desk | `/operate/visitors` | 1. Search visitor by mobile or scan pass.<br>2. Click 'Check In'. | Visitor marked 'Checked-In', live building occupancy increments, and host receives instant notification. |
| **TC-08**<br>Compliance Head | `/operate/compliance` | 1. Review 4 statutory categories.<br>2. Click 'Fire Safety NOC' card. | Shows validity dates, attached inspection certificate, renewal checklist, and 90-day warning badge. |
| **TC-09**<br>Prospect | `/` (Homepage)<br>`/marketplace` | 1. Click 'Book a Demo' or 'Schedule Tour'. | Calendly modal opens smoothly with `https://calendly.com/admin-cleanasset/30min`. Zero page reload. |
| **TC-10**<br>Super Admin | `/login` | 1. Enter `admin@officex.in`.<br>2. Submit credentials. | Enforces mandatory Level-2 MFA before granting access to the Super Admin Governance Console (`/admin`). |

---

## 9. Technical Handover & Production Configuration

For the client's IT and technical operations team, the following environment variables in `.env.local` govern core platform integrations. All services operate under 100% free production-ready tiers:

* `SMTP_HOST`: Standard SMTP server address (e.g. `'smtp.gmail.com'` or `'smtp-relay.brevo.com'`).
* `SMTP_PORT`: Secure port (`'465'` for SSL or `'587'` for TLS).
* `SMTP_USER`: Outgoing sender mailbox address.
* `SMTP_PASS`: 16-character application-specific password (Google App Password / Brevo API key).
* `SMTP_FROM`: Branded sender header (e.g. `'OfficeX Security <noreply@officex.in>'`).
* `NEXT_PUBLIC_SUPABASE_URL` & `ANON_KEY`: High-performance PostgreSQL database backend.

---
*Delivered & Certified by Scalezix Engineering & Delivery Team*  
*OfficeX Platform Ecosystem · Workspaces, Simplified · Confidential Document for Client UAT*
"""

md_path = os.path.abspath("client response/OfficeX_End_To_End_Client_Testing_Guide.md")
with open(md_path, "w", encoding="utf-8") as f:
    f.write(md_content)

print(f"✅ Successfully generated Markdown Document: {md_path}")
